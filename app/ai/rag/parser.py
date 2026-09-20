"""文档解析:从 PDF / DOCX / Markdown / 纯文本中抽取纯文本(第 14 阶段 M1)。

依赖 pymupdf(PDF)与 python-docx(DOCX)。两者都是纯 Python 侧封装,
不额外起服务,符合"M1 不引入额外中间件"的约束。

注意:扫描版 PDF(图片型)抽不出文字,这里不做 OCR —— 属明确不在 M1 范围内。
这类文件解析后会得到空文本,入库时会被标记为失败并给出原因。
"""
from pathlib import Path

from app.ai.rag.config import SUPPORTED_EXTENSIONS


class UnsupportedFileTypeError(ValueError):
    """文件类型不在支持列表内。"""


def is_supported(filename: str) -> bool:
    """扩展名是否受支持。"""
    return Path(filename).suffix.lower() in SUPPORTED_EXTENSIONS


def extract_text(file_path: str | Path) -> str:
    """按扩展名分派解析器,返回纯文本。

    抛出 UnsupportedFileTypeError 表示类型不支持;
    其他异常(文件损坏、加密等)原样向上抛,由调用方记录到 document.status。
    """
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"不支持的文件类型: {suffix}。当前支持 {sorted(SUPPORTED_EXTENSIONS)}"
        )

    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    return _extract_plain(path)


def _extract_pdf(path: Path) -> str:
    """用 pymupdf 逐页抽取文本。"""
    import fitz  # pymupdf

    parts: list[str] = []
    with fitz.open(path) as doc:
        for page in doc:
            text = page.get_text().strip()
            if text:
                parts.append(text)
    return "\n\n".join(parts)


def _extract_docx(path: Path) -> str:
    """用 python-docx 抽取段落与表格文本。"""
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # 表格内容也要带上 —— 产品文档里价格、时效这类关键信息常在表格里
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _extract_plain(path: Path) -> str:
    """纯文本 / Markdown:优先 UTF-8,失败则退回 GBK(中文 Windows 常见)。"""
    raw = path.read_bytes()
    for encoding in ("utf-8", "gbk"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")
